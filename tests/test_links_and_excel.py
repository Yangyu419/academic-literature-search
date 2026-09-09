from pathlib import Path

from openpyxl import load_workbook

from literature_finder.excel_writer import HEADERS, write_excel
from literature_finder.link_resolver import BestLegalAccessResolver
from literature_finder.models import LiteratureRecord


def test_arxiv_verified_download_is_preserved():
    record = LiteratureRecord("A preprint", open_access_url="https://arxiv.org/abs/1234.5678", publisher_url="https://arxiv.org/abs/1234.5678", download_url="https://arxiv.org/pdf/1234.5678", download_source="arXiv", download_file_type="pdf", download_permission_verified=True)
    BestLegalAccessResolver().resolve(record)
    assert record.best_access_url == "https://arxiv.org/abs/1234.5678"
    assert record.download_permission_verified is True


def test_excel_has_requested_columns_without_doi_or_link(tmp_path: Path):
    record = LiteratureRecord(
        "A fuel performance model",
        literature_type="期刊论文",
        publication_date="2024-01-02",
        source="Journal of Fuel Engineering",
        doi="10.1234/test",
        best_access_url="https://doi.org/10.1234/test",
        title_zh="燃料性能模型",
        research_content="建立燃料性能预测模型，并评估其在不同运行工况下的适用性。",
    )
    output = write_excel([record], tmp_path / "out.xlsx")
    sheet = load_workbook(output).active
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:H2"
    row_values = [cell.value for cell in sheet[2]]
    assert row_values[2] == "A fuel performance model"
    assert row_values[3] == "燃料性能模型"
    assert row_values[6] == "建立燃料性能预测模型，并评估其在不同运行工况下的适用性。"
    assert "10.1234/test" not in row_values
    assert "https://doi.org/10.1234/test" not in row_values
    assert all(cell.hyperlink is None for cell in sheet[2])


def test_excel_marks_missing_english_translation_and_abstract(tmp_path: Path):
    record = LiteratureRecord("An English title")
    output = write_excel([record], tmp_path / "out.xlsx")
    row = list(load_workbook(output).active[2])
    assert row[3].value == "未翻译，需人工补充"
    assert row[6].value == "未获取摘要，需人工补充"
