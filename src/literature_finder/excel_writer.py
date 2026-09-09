"""Simple compatible XLSX export with exactly the requested eight columns.

DOI and access URLs are deliberately kept as internal ``LiteratureRecord``
fields for the lawful-download workflow; they are not exposed in the sheet.
"""

from __future__ import annotations

from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import LiteratureRecord

HEADERS = ["序号", "文献类型", "文献名", "中文名", "文献发表日期", "期刊/来源", "研究内容", "备注"]


def _is_probably_english_title(title: str) -> bool:
    letters = [character for character in title if character.isalpha()]
    return bool(letters) and not any("\u4e00" <= character <= "\u9fff" for character in title)


def _display_chinese_title(record: LiteratureRecord) -> str:
    if record.title_zh:
        return record.title_zh
    if _is_probably_english_title(record.title):
        return "未翻译，需人工补充"
    return record.title


def _display_research_content(record: LiteratureRecord) -> str:
    if record.research_content:
        return record.research_content
    if record.abstract:
        # Keep a bounded, readable fallback without inventing content.  The
        # search agent should provide a concise summary when an abstract exists;
        # otherwise retain at most the first three abstract sentences.
        compact = " ".join(record.abstract.split())
        sentences = [part.strip() for part in re.split(r"(?<=[。！？.!?])\s*", compact) if part.strip()]
        summary = " ".join(sentences[:3]) if sentences else compact
        return summary if len(summary) <= 240 else summary[:237] + "..."
    return "未获取摘要，需人工补充"


def write_excel(records: list[LiteratureRecord], path: str | Path) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Literature"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, record in enumerate(records, 1):
        sheet.append([
            index,
            record.literature_type or "其他",
            record.title,
            _display_chinese_title(record),
            record.publication_date or "",
            record.source or "",
            _display_research_content(record),
            record.notes or "",
        ])
        sheet.cell(index + 1, 5).number_format = "@"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [8, 14, 52, 36, 18, 28, 60, 44]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(output)
    return output
